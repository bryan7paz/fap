"""Geração do relatório .docx de um repositório (score, métricas, gráficos).

Usado por GET /repo/<id>/relatorio. Os gráficos são PNGs em memória
(matplotlib Agg) embutidos no documento.
"""
import io
from datetime import datetime

import matplotlib

matplotlib.use("Agg")  # backend sem display — só geração de arquivo
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches


def _png_commits(serie):
    """Gráfico de linha: commits por dia. Retorna BytesIO ou None."""
    if not serie:
        return None
    dias = [p["dia"] for p in serie]
    commits = [p["commits"] for p in serie]
    fig, ax = plt.subplots(figsize=(7, 2.6))
    ax.plot(dias, commits, color="#0f766e", linewidth=1.6)
    ax.fill_between(dias, commits, color="#0f766e", alpha=0.10)
    ax.set_ylabel("commits")
    ax.grid(alpha=0.25)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def _png_curva(curva):
    """Gráfico de linha: participação top-3 e top-1 por mês. BytesIO ou None."""
    if not curva:
        return None
    meses = [c["mes"] for c in curva]
    fig, ax = plt.subplots(figsize=(7, 2.6))
    ax.plot(meses, [c["top3"] for c in curva], marker="o",
            color="#0f766e", linewidth=2.0, label="top-3")
    ax.plot(meses, [c["top1"] for c in curva], marker="s",
            color="#b45309", linewidth=1.6, linestyle="--", label="top-1")
    ax.set_ylabel("% dos commits do mês")
    ax.set_ylim(0, 105)
    ax.legend()
    ax.grid(alpha=0.25)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def gerar_relatorio(dados):
    """Monta o .docx a partir do pacote de dados do repositório; retorna BytesIO."""
    repo = dados["repo"]
    metricas = dados["metricas"]
    score = dados["score"]
    totais = dados["totais"]

    doc = Document()
    doc.add_heading(f"Relatório FAP — {repo['nome']}", 0)

    doc.add_paragraph(f"Repositório: {repo['url']}")
    if metricas:
        doc.add_paragraph(
            f"Período analisado: {metricas['periodo_inicio']} a {metricas['periodo_fim']}"
        )
    doc.add_paragraph(
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} pela "
        "Framework Analytics Platform (FAP)."
    )

    doc.add_heading("Score de sustentabilidade", level=1)
    if score["score"] is None:
        doc.add_paragraph("Sem dados suficientes para calcular o score.")
    else:
        doc.add_paragraph(f"Score composto (0–100): {score['score']}")
        tabela = doc.add_table(rows=1, cols=3)
        tabela.style = "Light Grid Accent 1"
        cab = tabela.rows[0].cells
        cab[0].text, cab[1].text, cab[2].text = "Componente", "Valor", "Observação"
        for comp in score["componentes"]:
            celulas = tabela.add_row().cells
            celulas[0].text = comp["nome"]
            celulas[1].text = f"{comp['valor']}"
            celulas[2].text = comp.get("alerta") or comp["descricao"]

    doc.add_heading("Métricas da janela", level=1)
    linhas = [
        ("Commits na janela", totais["commits"]),
        ("Linhas adicionadas", totais["linhas_add"]),
        ("Linhas removidas", totais["linhas_del"]),
    ]
    if metricas:
        linhas += [
            ("Bus Factor", metricas["bus_factor"]),
            ("TTFR (mediana, dias)", metricas["ttfr"]),
            ("Churn relativo", metricas["churn_relativo"]),
            ("Cadência (releases/mês)", metricas["cadencia_releases"]),
            ("Issues abertas", metricas["issues_abertas"]),
            ("Issues fechadas", metricas["issues_fechadas"]),
        ]
    tabela = doc.add_table(rows=1, cols=2)
    tabela.style = "Light Grid Accent 1"
    cab = tabela.rows[0].cells
    cab[0].text, cab[1].text = "Métrica", "Valor"
    for rotulo, valor in linhas:
        celulas = tabela.add_row().cells
        celulas[0].text = rotulo
        celulas[1].text = "—" if valor is None else str(valor)

    png = _png_commits(dados["serie"])
    if png:
        doc.add_heading("Atividade — commits por dia", level=1)
        doc.add_picture(png, width=Inches(6.5))

    png = _png_curva(dados["curva"])
    if png:
        doc.add_heading("Curva de concentração de conhecimento", level=1)
        doc.add_paragraph(
            "Participação do top-3 (e top-1) de autores nos commits de cada mês."
        )
        doc.add_picture(png, width=Inches(6.5))

    doc.add_paragraph(
        "Metodologia: Bus Factor = menor k com participação acumulada > 50%; "
        "TTFR = mediana do tempo até a primeira resposta humana (issues do "
        "período, sem PRs e sem bots); churn relativo = (linhas +/−) / LOC; "
        "score = média dos componentes normalizados (0–100)."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
