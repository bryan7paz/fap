"""Relatório .docx e snapshot avaliável — geração em memória e rotas."""
import io

import pytest

from relatorio import gerar_relatorio


def _dados_minimos():
    return {
        "repo": {"nome": "demo", "url": "https://github.com/a/b.git"},
        "metricas": None,
        "score": {"score": None, "componentes": []},
        "totais": {"commits": 0, "linhas_add": 0, "linhas_del": 0},
        "serie": [],
        "curva": [],
    }


def _dados_completos():
    return {
        "repo": {"nome": "demo", "url": "https://github.com/a/b.git"},
        "metricas": {
            "periodo_inicio": "2026-03-01", "periodo_fim": "2026-09-01",
            "bus_factor": 2, "ttfr": 1.5, "churn_relativo": 0.3,
            "cadencia_releases": 2.0, "issues_abertas": 5, "issues_fechadas": 10,
        },
        "score": {"score": 80.0, "componentes": [
            {"nome": "Atividade", "descricao": "x commits", "valor": 80.0},
            {"nome": "Bus Factor", "descricao": "BF = 2", "valor": 40.0,
             "alerta": "Conhecimento concentrado em 1–2 pessoas — risco."},
        ]},
        "totais": {"commits": 123, "linhas_add": 100, "linhas_del": 50},
        "serie": [{"dia": "2026-09-01", "commits": 3},
                  {"dia": "2026-09-02", "commits": 5}],
        "curva": [{"mes": "2026-08", "top3": 70.0, "top1": 40.0},
                  {"mes": "2026-09", "top3": 60.0, "top1": 30.0}],
    }


def test_relatorio_minimo_gera_docx_valido():
    buf = gerar_relatorio(_dados_minimos())
    assert buf.read(2) == b"PK"  # docx é um zip


def test_relatorio_completo_tem_titulo_e_alerta():
    buf = gerar_relatorio(_dados_completos())
    from docx import Document
    doc = Document(io.BytesIO(buf.getvalue()))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "Relatório FAP — demo" in texto
    tabelas = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "risco" in tabelas  # alerta do Bus Factor foi parar no documento


@pytest.fixture()
def client():
    import app as app_mod
    app_mod.app.config["TESTING"] = True
    with app_mod.app.test_client() as c:
        yield c


@pytest.fixture()
def repo_do_teste(client):
    from database import (desvincular_e_limpar, link_usuario_repositorio,
                          upsert_repositorio, upsert_usuario)
    id_usuario = upsert_usuario(github_id=987654321, login="dev-teste",
                                nome="Teste")
    id_repo = upsert_repositorio("relatorio-teste",
                                 "https://github.com/teste-fap/relatorio-teste.git")
    link_usuario_repositorio(id_usuario, id_repo)
    with client.session_transaction() as s:
        s["_user_id"] = str(id_usuario)
    yield id_repo
    desvincular_e_limpar(id_usuario, id_repo)


def test_relatorio_route_sem_login(client):
    assert client.get("/repo/1/relatorio").status_code == 302


def test_relatorio_route_owner_baixa_docx(client, repo_do_teste):
    resp = client.get(f"/repo/{repo_do_teste}/relatorio")
    assert resp.status_code == 200
    assert resp.get_data()[:2] == b"PK"


def test_snapshot_sem_login(client):
    assert client.get("/snapshot").status_code == 302
    assert client.get("/snapshot.csv?periodo=2026-01-01%7C2026-07-01").status_code == 302


def test_snapshot_logado_sem_periodos(client, repo_do_teste):
    resp = client.get("/snapshot")
    assert resp.status_code == 200
    assert "Nenhum período coletado" in resp.get_data(as_text=True)


def test_snapshot_periodo_invalido(client, repo_do_teste):
    resp = client.get("/snapshot?periodo=batata")
    assert resp.status_code == 200  # ignora o parâmetro e mostra o formulário


def test_snapshot_csv_sem_periodo_400(client, repo_do_teste):
    assert client.get("/snapshot.csv").status_code == 400


def test_snapshot_csv_com_periodo_valido(client, repo_do_teste):
    resp = client.get("/snapshot.csv?periodo=2026-01-01|2026-07-01")
    assert resp.status_code == 200
    texto = resp.get_data(as_text=True)
    assert texto.splitlines()[0].startswith("repositorio,")
